#!/usr/bin/env python3
"""
price_loader.py — загрузка и матчинг прайс-листа.

Предоставляет:
  _normalize_name(name)  → str           — нормализация для сравнения
  _parse_price(value)    → int | None    — парсинг числовой цены
  PriceIndex             — индекс с методами load() и match()

Пайплайн не падает при отсутствии/повреждении price.xlsx:
load() возвращает False и логирует warning.
"""

import heapq
import logging
import re
from difflib import SequenceMatcher

log = logging.getLogger(__name__)

try:
    import config as _cfg
    _PRICE_MATCH_THRESHOLD = getattr(_cfg, 'PRICE_MATCH_THRESHOLD', 0.85)
    _PRICE_AMBIGUITY_GAP = getattr(_cfg, 'PRICE_AMBIGUITY_GAP', 0.02)
except ImportError:
    _PRICE_MATCH_THRESHOLD = 0.85
    _PRICE_AMBIGUITY_GAP = 0.02

_CURRENCY_WORDS_RE = re.compile(r'\b(?:ТГ|Т|KZT)\b', re.UNICODE)
_PUNCT_RE = re.compile(r'[^\w\s]', re.UNICODE)
_SPACES_RE = re.compile(r'\s+')
_NON_DIGIT_RE = re.compile(r'\D')


def _normalize_name(name: str) -> str:
    """Нормализует название для индексирования/матчинга."""
    if not name:
        return ""
    s = str(name).upper().strip()
    # Ё → Е
    s = s.replace('Ё', 'Е')
    # Убрать символ валюты ₸
    s = s.replace('₸', '')
    # Убрать слова-валюты (ТГ, Т, KZT) по границам слов
    s = _CURRENCY_WORDS_RE.sub('', s)
    # Убрать пунктуацию (оставить \w и пробелы)
    s = _PUNCT_RE.sub(' ', s)
    # Collapse spaces
    s = _SPACES_RE.sub(' ', s).strip()
    return s


def _parse_price(value) -> 'int | None':
    """Парсит значение цены в int или возвращает None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        v = int(value)
        return v if v > 0 else None
    if isinstance(value, str):
        digits = _NON_DIGIT_RE.sub('', value)
        if not digits:
            return None
        v = int(digits)
        return v if v > 0 else None
    return None


class PriceIndex:
    """
    Индекс прайс-листа.

    Загружается из price.xlsx (лист «Стоимость»,
    колонки «Название» / «Стоимость»).
    При ошибке/отсутствии файла — мягкое warning, пайплайн не падает.
    """

    def __init__(self):
        self._index: dict = {}   # normalized_name → price
        self._loaded: bool = False
        self._fuzzy_cache: dict = {}  # normalized_query → (price, score, ambiguous)

    def load(self, path) -> bool:
        """
        Загружает прайс-лист из файла.
        Returns True при успехе, False при любой ошибке.
        """
        try:
            import openpyxl
        except ImportError:
            log.warning("openpyxl не установлен — прайс-лист не загружен")
            return False

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            log.warning(f"Прайс-лист не найден: {path}")
            return False
        except Exception as e:
            log.warning(f"Не удалось открыть прайс-лист {path}: {e}")
            return False

        try:
            if "Стоимость" not in wb.sheetnames:
                log.warning(
                    f"Лист 'Стоимость' не найден в {path} "
                    f"(доступные: {list(wb.sheetnames)})"
                )
                wb.close()
                return False

            ws = wb["Стоимость"]
            index: dict = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 2:
                    continue
                name_raw, price_raw = row[0], row[1]
                if name_raw is None:
                    continue
                norm = _normalize_name(str(name_raw))
                if not norm:
                    continue
                price = _parse_price(price_raw)
                if price is None:
                    continue
                if norm in index:
                    log.warning(
                        f"Дубль в прайсе: '{norm}' "
                        f"(из '{name_raw}') — первое вхождение сохранено"
                    )
                    continue
                index[norm] = price

            self._index = index
            self._loaded = True
            self._fuzzy_cache = {}
            log.info(f"Прайс-лист загружен: {len(index)} позиций из {path}")
            wb.close()
            return True

        except Exception as e:
            log.warning(f"Ошибка при чтении прайс-листа: {e}")
            try:
                wb.close()
            except Exception:
                pass
            return False

    @property
    def is_loaded(self) -> bool:
        return self._loaded

    def match(self, name: str) -> tuple:
        """
        Находит цену для позиции по имени.

        Returns:
            (price, score, ambiguous) где:
            - price     — цена из прайса (int) или None
            - score     — степень совпадения (float 0-1) или None
            - ambiguous — True если два лучших варианта слишком близки (diff < 0.02)
        """
        if not self._loaded or not self._index:
            return (None, None, False)

        norm = _normalize_name(name)
        if not norm:
            return (None, None, False)

        # 1. Точное совпадение
        if norm in self._index:
            return (self._index[norm], 1.0, False)

        # 2. Нечёткий поиск (с кэшем)
        if norm in self._fuzzy_cache:
            return self._fuzzy_cache[norm]

        keys = list(self._index.keys())
        if not keys:
            return (None, None, False)

        top2 = heapq.nlargest(
            2,
            ((SequenceMatcher(None, norm, k).ratio(), k) for k in keys),
            key=lambda x: x[0],
        )
        top_score, top_key = top2[0]

        if top_score < _PRICE_MATCH_THRESHOLD:
            result = (None, top_score, False)
            self._fuzzy_cache[norm] = result
            return result

        # Проверка неоднозначности
        if len(top2) >= 2:
            second_score = top2[1][0]
            if abs(top_score - second_score) < _PRICE_AMBIGUITY_GAP:
                log.warning(
                    f"Неоднозначное совпадение для '{name}': "
                    f"'{top_key}'={top_score:.3f} vs "
                    f"'{top2[1][1]}'={second_score:.3f}"
                )
                result = (None, top_score, True)
                self._fuzzy_cache[norm] = result
                return result

        result = (self._index[top_key], top_score, False)
        self._fuzzy_cache[norm] = result
        return result
