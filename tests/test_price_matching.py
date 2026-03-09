"""
Тесты модуля price_loader и интеграции прайс-листа в пайплайн.

Охватывают:
  - _normalize_name / _parse_price
  - PriceIndex.match (exact, fuzzy, ambiguous, not-loaded)
  - логику подстановки/сравнения цен
  - write_to_excel — новые колонки *_прайс / *_дельта / *_совпадает
  - generate_pipeline_report — лист «Прайс-сверка»
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from price_loader import PriceIndex, _normalize_name, _parse_price


# ─────────────────────────────────────────────────────────────────────────────
# HELPER
# ─────────────────────────────────────────────────────────────────────────────

def _make_price_index(data: dict) -> PriceIndex:
    """Создаёт PriceIndex с известными данными, минуя загрузку файла."""
    pi = PriceIndex()
    pi._index = {_normalize_name(k): v for k, v in data.items()}
    pi._loaded = True
    return pi


# ─────────────────────────────────────────────────────────────────────────────
# TestNormalizeName
# ─────────────────────────────────────────────────────────────────────────────

class TestNormalizeName:

    def test_trim_upper(self):
        assert _normalize_name(" Массаж лица ") == "МАССАЖ ЛИЦА"

    def test_yo_to_ye(self):
        assert _normalize_name("ОБЁРТЫВАНИЕ") == "ОБЕРТЫВАНИЕ"

    def test_remove_currency(self):
        result = _normalize_name("ЛАЗЕР, тг")
        assert result == "ЛАЗЕР"


# ─────────────────────────────────────────────────────────────────────────────
# TestParsePrice
# ─────────────────────────────────────────────────────────────────────────────

class TestParsePrice:

    def test_string_spaces(self):
        assert _parse_price("5 000 тг") == 5000

    def test_numeric(self):
        assert _parse_price(5000) == 5000

    def test_no_digits(self):
        assert _parse_price("тг") is None

    def test_none(self):
        assert _parse_price(None) is None


# ─────────────────────────────────────────────────────────────────────────────
# TestPriceIndexMatch
# ─────────────────────────────────────────────────────────────────────────────

class TestPriceIndexMatch:

    def test_exact(self):
        pi = _make_price_index({"Массаж лица": 5000})
        price, score, ambiguous = pi.match("Массаж лица")
        assert price == 5000
        assert score == 1.0
        assert ambiguous is False

    def test_below_threshold(self):
        pi = _make_price_index({"Массаж лица": 5000})
        price, score, ambiguous = pi.match("Абсолютно совершенно иная вещь xyz")
        assert price is None

    def test_ambiguous(self):
        pi = PriceIndex()
        # Два похожих названия: SequenceMatcher даст близкие scores → ambiguous
        pi._index = {"МАССАЖ ЛИЦА А": 5000, "МАССАЖ ЛИЦА Б": 6000}
        pi._loaded = True
        price, score, ambiguous = pi.match("Массаж лица")
        assert ambiguous is True
        assert price is None

    def test_not_loaded(self):
        pi = PriceIndex()
        price, score, ambiguous = pi.match("Массаж лица")
        assert price is None
        assert score is None
        assert ambiguous is False


# ─────────────────────────────────────────────────────────────────────────────
# TestPriceFill — логика подстановки/сравнения
# ─────────────────────────────────────────────────────────────────────────────

class TestPriceFill:

    def test_fill_empty_cost(self):
        """Пустая OCR-цена + найден прайс → подставляется."""
        pi = _make_price_index({"Массаж лица": 5000})
        p_price, _score, p_amb = pi.match("Массаж лица")
        ocr_val = _parse_price(None)

        if ocr_val is None and p_price is not None and not p_amb:
            final_val = p_price
            filled = True
        else:
            final_val = ocr_val
            filled = False

        assert final_val == 5000
        assert filled is True

    def test_mismatch_flag(self):
        """OCR-цена отличается >1% и >100 → совпадает=False."""
        pi = _make_price_index({"Массаж лица": 5000})
        p_price, _, _ = pi.match("Массаж лица")
        final_val = 9999  # delta = 4999, > max(50, 100)
        delta = final_val - p_price
        matches = abs(delta) <= max(0.01 * p_price, 100)
        assert matches is False

    def test_within_tolerance(self):
        """Отклонение ≤ max(1%, 100тг) → совпадает=True."""
        pi = _make_price_index({"Массаж лица": 5000})
        p_price, _, _ = pi.match("Массаж лица")
        final_val = 5050  # delta = 50, max(50, 100) = 100 → 50 ≤ 100
        delta = final_val - p_price
        matches = abs(delta) <= max(0.01 * p_price, 100)
        assert matches is True

    def test_ambiguous_no_fill(self):
        """Неоднозначное совпадение → цена не подставляется."""
        pi = PriceIndex()
        pi._index = {"МАССАЖ ЛИЦА А": 5000, "МАССАЖ ЛИЦА Б": 6000}
        pi._loaded = True
        p_price, _, p_amb = pi.match("Массаж лица")
        assert p_amb is True

        ocr_val = None
        if ocr_val is None and p_price is not None and not p_amb:
            final_val = p_price
        else:
            final_val = ocr_val
        assert final_val is None


# ─────────────────────────────────────────────────────────────────────────────
# TestIntegration
# ─────────────────────────────────────────────────────────────────────────────

class TestIntegration:

    def test_no_price_file(self):
        """Несуществующий путь → is_loaded=False, без исключения."""
        pi = PriceIndex()
        result = pi.load("/nonexistent/path/price.xlsx")
        assert result is False
        assert pi.is_loaded is False

    def test_write_to_excel_price_cols(self, tmp_path, monkeypatch):
        """
        Синтетический grouped_clients (3 строки: пустая/совпадение/расхождение)
        → колонки Стоимость_прайс, _дельта, _совпадает корректны.
        """
        import config
        monkeypatch.setattr(config, "OUTPUT_FILE", str(tmp_path / "test_clients.xlsx"))

        from client_card_ocr import write_to_excel

        pi = _make_price_index({"Массаж лица": 5000})

        grouped = {
            "client_a": {
                "name": "Тест Клиент",
                "iin": "",
                "phone": "",
                "pages": [{
                    "page_type": "procedure_sheet",
                    "filename": "test.jpg",
                    "data": {
                        "procedures": [
                            {   # пустая цена → заполняется из прайса
                                "date": "01.01.2025",
                                "procedure_name": "Массаж лица",
                                "description": "",
                                "cost": None,
                            },
                            {   # цена совпадает → delta=0, совпадает=True
                                "date": "02.01.2025",
                                "procedure_name": "Массаж лица",
                                "description": "",
                                "cost": "5000",
                            },
                            {   # цена отличается → совпадает=False
                                "date": "03.01.2025",
                                "procedure_name": "Массаж лица",
                                "description": "",
                                "cost": "9999",
                            },
                        ]
                    }
                }]
            }
        }

        stats = write_to_excel(grouped, [], price_index=pi)

        assert stats["loaded"] is True
        assert stats["Процедуры"]["filled"] >= 1
        assert stats["Процедуры"]["mismatched"] >= 1

        import openpyxl
        wb = openpyxl.load_workbook(str(tmp_path / "test_clients.xlsx"))
        assert "Процедуры" in wb.sheetnames

        ws = wb["Процедуры"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Стоимость_прайс" in headers
        assert "Стоимость_дельта" in headers
        assert "Стоимость_совпадает" in headers

        prais_idx = headers.index("Стоимость_прайс") + 1
        delta_idx = headers.index("Стоимость_дельта") + 1
        match_idx = headers.index("Стоимость_совпадает") + 1
        cost_idx = headers.index("Стоимость") + 1

        # Строка 2: пустая цена → заполнена прайсом (5000)
        assert ws.cell(row=2, column=prais_idx).value == 5000
        assert ws.cell(row=2, column=cost_idx).value == 5000

        # Строка 3: cost=5000, совпадает → delta=0, совпадает=True
        assert ws.cell(row=3, column=delta_idx).value == 0
        assert ws.cell(row=3, column=match_idx).value is True

        # Строка 4: cost=9999, не совпадает → совпадает=False
        assert ws.cell(row=4, column=match_idx).value is False

        wb.close()

    def test_write_to_excel_append_price_cols(self, tmp_path, monkeypatch):
        """
        Append mode: файл существует → _append_new_clients применяет прайс
        для новых клиентов → price_stats.loaded=True, колонки добавлены.
        """
        import config
        monkeypatch.setattr(config, "OUTPUT_FILE", str(tmp_path / "test.xlsx"))

        from client_card_ocr import write_to_excel

        pi = _make_price_index({"Массаж лица": 5000})

        # Шаг 1: создать файл БЕЗ прайса (exists → append path next time)
        grouped_existing = {
            "client_a": {
                "name": "Старый Клиент",
                "iin": "",
                "phone": "",
                "pages": [{
                    "page_type": "procedure_sheet",
                    "filename": "old.jpg",
                    "data": {"procedures": [
                        {"date": "01.01.2025", "procedure_name": "Массаж лица",
                         "description": "", "cost": "5000"},
                    ]}
                }]
            }
        }
        stats1 = write_to_excel(grouped_existing, [], price_index=None)
        assert stats1["loaded"] is False  # создан без прайса

        # Шаг 2: дозапись с новым клиентом + price_index
        grouped_with_new = dict(grouped_existing)
        grouped_with_new["client_b"] = {
            "name": "Новый Клиент",
            "iin": "",
            "phone": "",
            "pages": [{
                "page_type": "procedure_sheet",
                "filename": "new.jpg",   # новый файл → новый клиент
                "data": {"procedures": [
                    {"date": "02.01.2025", "procedure_name": "Массаж лица",
                     "description": "", "cost": None},   # пустая → заполнить
                ]}
            }]
        }
        stats2 = write_to_excel(grouped_with_new, [], price_index=pi)

        assert stats2["loaded"] is True
        assert stats2["Процедуры"]["filled"] >= 1   # пустая цена заполнена

        import openpyxl
        wb = openpyxl.load_workbook(str(tmp_path / "test.xlsx"))
        ws = wb["Процедуры"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Стоимость_прайс" in headers
        assert "Стоимость_дельта" in headers
        assert "Стоимость_совпадает" in headers

        # Строка с новым клиентом (последняя) должна иметь filled cost=5000
        prais_idx = headers.index("Стоимость_прайс") + 1
        cost_idx  = headers.index("Стоимость") + 1
        last_row = ws.max_row
        assert ws.cell(row=last_row, column=prais_idx).value == 5000
        assert ws.cell(row=last_row, column=cost_idx).value == 5000   # заполнено
        wb.close()

    def test_pipeline_report_section(self, tmp_path, monkeypatch):
        """price_stats с данными → лист «Прайс-сверка» присутствует в отчёте."""
        import logging
        import config
        import run_pipeline

        monkeypatch.setattr(run_pipeline, "__file__", str(tmp_path / "run_pipeline.py"))

        log = logging.getLogger("test_pipeline_report")
        price_stats = {
            "loaded": True,
            "Процедуры": {"found": 10, "filled": 3, "mismatched": 1, "ambiguous": 0},
            "Покупки":   {"found": 5,  "filled": 1, "mismatched": 0, "ambiguous": 1},
            "Комплексы": {"found": 2,  "filled": 0, "mismatched": 0, "ambiguous": 0},
        }

        result = run_pipeline.generate_pipeline_report(
            log, config, None, None, price_stats=price_stats
        )

        assert result is not None
        import openpyxl
        wb = openpyxl.load_workbook(result)
        assert "Прайс-сверка" in wb.sheetnames

        # Проверяем, что в листе есть строки для всех трёх листов
        ws = wb["Прайс-сверка"]
        sheet_col = 1
        values = [ws.cell(row=r, column=sheet_col).value for r in range(2, ws.max_row + 1)]
        assert "Процедуры" in values
        assert "Покупки" in values
        assert "Комплексы" in values
        wb.close()

    def test_pipeline_report_no_price(self, tmp_path, monkeypatch):
        """price_stats=None → отчёт генерируется без ошибок, лист «Прайс-сверка» есть."""
        import logging
        import config
        import run_pipeline

        monkeypatch.setattr(run_pipeline, "__file__", str(tmp_path / "run_pipeline.py"))

        log = logging.getLogger("test_pipeline_report_no_price")

        result = run_pipeline.generate_pipeline_report(
            log, config, None, None, price_stats=None
        )

        assert result is not None
        import openpyxl
        wb = openpyxl.load_workbook(result)
        assert "Прайс-сверка" in wb.sheetnames
        wb.close()
