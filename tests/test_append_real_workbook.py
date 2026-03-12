"""
Smoke test: append-mode on real clients_database.xlsx.

Copies the real workbook, appends a synthetic new client with
procedures/purchases/complexes using exact names from price.xlsx,
and verifies that:
  - existing rows are preserved cell-by-cell
  - new rows are appended correctly
  - price matching fills expected columns

All expectations are computed from the before-snapshot so the test
stays green as the real workbook grows.
"""
import re
import shutil
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
REAL_WORKBOOK = ROOT / "clients_database.xlsx"
REAL_PRICE = ROOT / "data" / "price.xlsx"

def _workbook_readable(path):
    """Check if an xlsx file can actually be opened (not just exists)."""
    if not path.exists():
        return False
    try:
        wb = openpyxl.load_workbook(path)
        wb.close()
        return True
    except Exception:
        return False


# Skip entire module if real files are missing or unreadable (e.g. CI without data,
# or workbook corrupted by interrupted write)
pytestmark = pytest.mark.skipif(
    not _workbook_readable(REAL_WORKBOOK) or not REAL_PRICE.exists(),
    reason="Real clients_database.xlsx unreadable/missing or price.xlsx not found",
)

# --- Exact names from price.xlsx (rows 2, 3, 10) ---
PRICE_NAME_1 = "КОНСУЛЬТАЦИЯ КОСМЕТОЛОГА подбор домашнего ухода и/или процедур (30 мин)"
PRICE_COST_1 = 5000
PRICE_NAME_2 = "КОНСУЛЬТАЦИЯ ДЕРМАТОЛОГА ПО АКНЕ И РОЗАЦИИ"
PRICE_COST_2 = 15000
PRICE_NAME_3 = "ЧИСТКА ЛИЦА КЛАССИЧЕСКАЯ"
PRICE_COST_3 = 20000

# Unique filename that guarantees the client is treated as new
SMOKE_FILENAME = "smoke_append_test_001.jpg"


def _snapshot_workbook(path):
    """Return {sheet_name: {"rows": [[cell_values...]], "max_row": int, "max_col": int}}."""
    wb = openpyxl.load_workbook(path)
    snapshot = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for r in range(1, ws.max_row + 1):
            row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            rows.append(row)
        snapshot[name] = {
            "rows": rows,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
        }
    wb.close()
    return snapshot


def _max_cl_id(snapshot):
    """Extract the largest CL-NNNN number from the Клиенты sheet."""
    max_num = 0
    for row in snapshot["Клиенты"]["rows"][1:]:  # skip header
        if row and row[0]:
            m = re.match(r"CL-(\d+)", str(row[0]))
            if m:
                max_num = max(max_num, int(m.group(1)))
    return max_num


def _assert_existing_rows_unchanged(ws, before_sheet, sheet_name):
    """Cell-by-cell comparison of all rows that existed before the append."""
    for r_idx, before_row in enumerate(before_sheet["rows"]):
        for c_idx, expected_val in enumerate(before_row):
            actual = ws.cell(r_idx + 1, c_idx + 1).value
            assert actual == expected_val, (
                f"{sheet_name}[{r_idx+1},{c_idx+1}] changed: "
                f"{expected_val!r} → {actual!r}"
            )


def _build_synthetic_client():
    """Build a grouped_clients dict with one new client."""
    return {
        "smoke_test_client": {
            "name": "Тестовый Клиент Дымовой",
            "iin": "",
            "phone": "+70001112233",
            "pages": [
                {
                    "page_type": "medical_card_front",
                    "filename": SMOKE_FILENAME,
                    "data": {},
                },
                {
                    "page_type": "procedure_sheet",
                    "filename": SMOKE_FILENAME,
                    "data": {
                        "procedures": [
                            {
                                "date": "01.03.2026",
                                "procedure_name": PRICE_NAME_1,
                                "description": "тест: пустая стоимость → filled",
                                "cost": None,
                            },
                            {
                                "date": "02.03.2026",
                                "procedure_name": PRICE_NAME_2,
                                "description": "тест: цена совпадает",
                                "cost": PRICE_COST_2,
                            },
                            {
                                "date": "03.03.2026",
                                "procedure_name": PRICE_NAME_3,
                                "description": "тест: цена не совпадает",
                                "cost": 99999,
                            },
                        ]
                    },
                },
                {
                    "page_type": "products_list",
                    "filename": SMOKE_FILENAME,
                    "data": {
                        "products": [
                            {
                                "date": "04.03.2026",
                                "consultant": "Тестовый врач",
                                "product_name": PRICE_NAME_1,
                                "price": PRICE_COST_1,
                            }
                        ]
                    },
                },
                {
                    "page_type": "complex_package",
                    "filename": SMOKE_FILENAME,
                    "data": {
                        "complex_name": PRICE_NAME_3,
                        "complex_cost": 25000,
                        "patient_name": "Тестовый Клиент Дымовой",
                        "contacts": "+70001112233",
                        "doctor": "Тест Доктор",
                        "purchase_date": "05.03.2026",
                        "procedures": [
                            {
                                "number": "1",
                                "procedure": "Процедура А",
                                "date": "06.03.2026",
                                "quantity": "1",
                                "comment": "",
                            }
                        ],
                    },
                },
            ],
        }
    }


@pytest.fixture()
def smoke_workbook(tmp_path, monkeypatch):
    """Copy real workbook to tmp_path, patch config, load PriceIndex."""
    import config
    from price_loader import PriceIndex

    dest = tmp_path / "clients_database.xlsx"
    shutil.copy2(REAL_WORKBOOK, dest)
    monkeypatch.setattr(config, "OUTPUT_FILE", str(dest))

    pi = PriceIndex()
    loaded = pi.load(str(REAL_PRICE))
    assert loaded, f"PriceIndex failed to load from {REAL_PRICE}"

    before = _snapshot_workbook(dest)
    return dest, pi, before


class TestAppendRealWorkbook:
    """Append a synthetic client to the real workbook copy."""

    def test_existing_rows_preserved_and_new_client_appended(self, smoke_workbook):
        from client_card_ocr import write_to_excel

        dest, pi, before = smoke_workbook
        grouped = _build_synthetic_client()

        # Derive expectations from the snapshot
        prev_max_id = _max_cl_id(before)
        expected_new_id = f"CL-{prev_max_id + 1:04d}"

        before_cl_rows = before["Клиенты"]["max_row"]
        before_proc_rows = before["Процедуры"]["max_row"]
        before_purch_rows = before["Покупки"]["max_row"]
        before_comp_rows = before["Комплексы"]["max_row"]

        stats = write_to_excel(grouped, [], price_index=pi)

        assert stats["loaded"] is True

        wb = openpyxl.load_workbook(dest)

        # --- Клиенты: +1 row ---
        ws_cl = wb["Клиенты"]
        assert ws_cl.max_row == before_cl_rows + 1, (
            f"Клиенты: expected {before_cl_rows}+1 rows, got {ws_cl.max_row}"
        )
        _assert_existing_rows_unchanged(ws_cl, before["Клиенты"], "Клиенты")

        new_cl_row = before_cl_rows + 1
        assert ws_cl.cell(new_cl_row, 1).value == expected_new_id
        assert "Тестовый Клиент Дымовой" in str(ws_cl.cell(new_cl_row, 4).value)

        # --- Процедуры: +3 rows (3 procedures) ---
        ws_proc = wb["Процедуры"]
        assert ws_proc.max_row == before_proc_rows + 3, (
            f"Процедуры: expected {before_proc_rows}+3 rows, got {ws_proc.max_row}"
        )
        _assert_existing_rows_unchanged(ws_proc, before["Процедуры"], "Процедуры")

        headers_proc = [ws_proc.cell(1, c).value for c in range(1, ws_proc.max_column + 1)]
        cost_col = headers_proc.index("Стоимость") + 1
        price_col = headers_proc.index("Стоимость_прайс") + 1
        delta_col = headers_proc.index("Стоимость_дельта") + 1
        match_col = headers_proc.index("Стоимость_совпадает") + 1

        # First new row: empty cost → filled from price
        r1 = before_proc_rows + 1
        assert ws_proc.cell(r1, 1).value == expected_new_id
        assert ws_proc.cell(r1, cost_col).value == PRICE_COST_1, "Empty cost should be filled"
        assert ws_proc.cell(r1, price_col).value == PRICE_COST_1
        assert ws_proc.cell(r1, delta_col).value == 0
        assert ws_proc.cell(r1, match_col).value is True

        # Second new row: matching cost
        r2 = before_proc_rows + 2
        assert ws_proc.cell(r2, cost_col).value == PRICE_COST_2
        assert ws_proc.cell(r2, price_col).value == PRICE_COST_2
        assert ws_proc.cell(r2, delta_col).value == 0
        assert ws_proc.cell(r2, match_col).value is True

        # Third new row: mismatched cost (99999 vs 20000)
        r3 = before_proc_rows + 3
        assert ws_proc.cell(r3, cost_col).value == 99999
        assert ws_proc.cell(r3, price_col).value == PRICE_COST_3
        assert ws_proc.cell(r3, delta_col).value == 99999 - PRICE_COST_3
        assert ws_proc.cell(r3, match_col).value is False

        # --- Покупки: +1 row ---
        ws_purch = wb["Покупки"]
        assert ws_purch.max_row == before_purch_rows + 1, (
            f"Покупки: expected {before_purch_rows}+1 rows, got {ws_purch.max_row}"
        )
        _assert_existing_rows_unchanged(ws_purch, before["Покупки"], "Покупки")
        new_purch_row = before_purch_rows + 1
        assert ws_purch.cell(new_purch_row, 1).value == expected_new_id

        headers_purch = [ws_purch.cell(1, c).value for c in range(1, ws_purch.max_column + 1)]
        p_price_col = headers_purch.index("Цена_прайс") + 1
        assert ws_purch.cell(new_purch_row, p_price_col).value == PRICE_COST_1

        # --- Комплексы: +1 row (1 complex procedure) ---
        ws_comp = wb["Комплексы"]
        assert ws_comp.max_row == before_comp_rows + 1, (
            f"Комплексы: expected {before_comp_rows}+1 rows, got {ws_comp.max_row}"
        )
        _assert_existing_rows_unchanged(ws_comp, before["Комплексы"], "Комплексы")

        new_comp_row = before_comp_rows + 1
        assert ws_comp.cell(new_comp_row, 1).value == expected_new_id

        headers_comp = [ws_comp.cell(1, c).value for c in range(1, ws_comp.max_column + 1)]
        comp_price_col = headers_comp.index("Стоимость_прайс") + 1
        assert ws_comp.cell(new_comp_row, comp_price_col).value == PRICE_COST_3

        # --- Сверка_БД: existing rows preserved ---
        _assert_existing_rows_unchanged(wb["Сверка_БД"], before["Сверка_БД"], "Сверка_БД")

        wb.close()

    def test_price_stats_counts(self, smoke_workbook):
        from client_card_ocr import write_to_excel

        dest, pi, _ = smoke_workbook
        grouped = _build_synthetic_client()
        stats = write_to_excel(grouped, [], price_index=pi)

        assert stats["loaded"] is True

        # Процедуры: 3 found (all matched in price), 1 filled (empty cost),
        # 1 mismatched (99999 vs 20000)
        assert stats["Процедуры"]["found"] == 3
        assert stats["Процедуры"]["filled"] == 1
        assert stats["Процедуры"]["mismatched"] == 1

        # Покупки: 1 found, 0 filled (cost was provided and matches)
        assert stats["Покупки"]["found"] == 1
        assert stats["Покупки"]["filled"] == 0
        assert stats["Покупки"]["mismatched"] == 0

        # Комплексы: 1 found (complex_cost=25000 vs price=20000 → mismatched)
        assert stats["Комплексы"]["found"] == 1
        assert stats["Комплексы"]["mismatched"] == 1

    def test_no_duplicate_on_rerun(self, smoke_workbook):
        """Running append twice with same filename should not duplicate rows."""
        from client_card_ocr import write_to_excel

        dest, pi, before = smoke_workbook
        grouped = _build_synthetic_client()

        before_cl_rows = before["Клиенты"]["max_row"]
        before_proc_rows = before["Процедуры"]["max_row"]

        write_to_excel(grouped, [], price_index=pi)
        stats2 = write_to_excel(grouped, [], price_index=pi)

        # Second run should detect existing files and add nothing
        wb = openpyxl.load_workbook(dest)
        assert wb["Клиенты"].max_row == before_cl_rows + 1, "Should not duplicate client"
        assert wb["Процедуры"].max_row == before_proc_rows + 3, "Should not duplicate procedures"
        wb.close()

        # price_stats should show zeros (no new rows written)
        assert stats2["Процедуры"]["found"] == 0
        assert stats2["Процедуры"]["filled"] == 0
